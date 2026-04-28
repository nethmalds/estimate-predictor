"""Excel report generator using openpyxl.

Produces a three-sheet workbook:
  Sheet 1 — Summary      : project parameters + total cost + confidence
  Sheet 2 — BOQ          : full item list with quantities, rates, costs, BSR codes
  Sheet 3 — Cost Breakdown: subtotals by category
"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from core.logging.logger import get_logger

logger = get_logger(__name__)

# ── Style constants ────────────────────────────────────────────────────────────
_DARK_FILL   = PatternFill("solid", fgColor="1F3864")
_MID_FILL    = PatternFill("solid", fgColor="2F5496")
_CAT_FILL    = PatternFill("solid", fgColor="D6E4F0")
_TOTAL_FILL  = PatternFill("solid", fgColor="FFF2CC")
_WHITE_FILL  = PatternFill("solid", fgColor="FFFFFF")

_HDR_FONT    = Font(bold=True, color="FFFFFF", size=11)
_SUBHDR_FONT = Font(bold=True, color="FFFFFF", size=10)
_CAT_FONT    = Font(bold=True, color="1F3864", size=10)
_TOTAL_FONT  = Font(bold=True, size=11)
_BODY_FONT   = Font(size=10)

_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
_RIGHT  = Alignment(horizontal="right",  vertical="center")

_THIN = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


# ── Public API ─────────────────────────────────────────────────────────────────

def generate_excel_report(report_payload: dict) -> bytes:
    """Build an Excel workbook from the pipeline's report payload and return raw bytes."""
    wb = Workbook()
    wb.remove(wb.active)  # remove default blank sheet

    _sheet_summary(wb, report_payload)
    _sheet_boq(wb, report_payload)
    _sheet_breakdown(wb, report_payload)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    logger.info("excel_report_generated sheets=%d", len(wb.sheetnames))
    return buf.getvalue()


# ── Sheet 1: Summary ──────────────────────────────────────────────────────────

def _sheet_summary(wb: Workbook, payload: dict) -> None:
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 42

    project_info: dict = payload.get("project_info") or {}
    parameters:   dict = project_info.get("parameters") or {}
    costs:        dict = payload.get("costs") or {}
    confidence:   dict = payload.get("confidence") or {}
    boq_items:    list = payload.get("boq_items") or []

    # Title
    ws.merge_cells("A1:B1")
    _style(ws["A1"], "CONSTRUCTION COST ESTIMATION REPORT",
           font=Font(bold=True, size=14, color="FFFFFF"),
           fill=_DARK_FILL, align=_CENTER, height=32, ws=ws, row=1)

    sections: list[tuple[str | None, Any]] = [
        ("PROJECT DETAILS", None),
        ("Building Type",      project_info.get("building_type") or "—"),
        ("Number of Floors",   project_info.get("floors") or "—"),
        ("Bedrooms",           parameters.get("bedrooms") or "—"),
        ("Bathrooms",          parameters.get("bathrooms") or "—"),
        ("Built-up Area",      parameters.get("built_up_area") or "—"),
        ("Finish Level",       parameters.get("finish_level") or "—"),
        ("Roof Type",          parameters.get("roof_type") or "—"),
        ("Ceiling Type",       parameters.get("ceiling_type") or "—"),
        (None, None),
        ("COST SUMMARY", None),
        ("Base Total (LKR)",   costs.get("base_total") or 0.0),
        ("Contingencies (5%)", costs.get("contingencies") or 0.0),
        ("Grand Total (LKR)",  costs.get("total") or 0.0),
        (None, None),
        ("ESTIMATE QUALITY", None),
        ("Confidence Score",   f"{(confidence.get('score') or 0) * 100:.1f}%"),
        ("Total BOQ Items",    len(boq_items)),
    ]

    for r, (label, value) in enumerate(sections, start=2):
        if label is None:
            continue
        ws.row_dimensions[r].height = 20
        a = ws.cell(row=r, column=1, value=label)
        b = ws.cell(row=r, column=2, value=value)
        a.border = _THIN
        b.border = _THIN
        a.font = _BODY_FONT
        b.font = _BODY_FONT
        a.alignment = _LEFT
        b.alignment = _LEFT

        if value is None:  # section header
            a.font = _SUBHDR_FONT
            a.fill = _MID_FILL
            b.fill = _MID_FILL
            ws.merge_cells(f"A{r}:B{r}")
        elif label == "Grand Total (LKR)":
            a.font = _TOTAL_FONT
            b.font = _TOTAL_FONT
            a.fill = _TOTAL_FILL
            b.fill = _TOTAL_FILL
            if isinstance(value, (int, float)):
                b.number_format = "#,##0.00"
                b.alignment = _RIGHT
        elif isinstance(value, (int, float)):
            b.number_format = "#,##0.00"
            b.alignment = _RIGHT


# ── Sheet 2: Bill of Quantities ───────────────────────────────────────────────

def _sheet_boq(wb: Workbook, payload: dict) -> None:
    ws = wb.create_sheet("Bill of Quantities")

    col_widths = [6, 48, 26, 10, 14, 16, 16, 18, 10]
    col_headers = ["No.", "Description", "Category", "Unit",
                   "Quantity", "Rate (LKR)", "Cost (LKR)", "BSR Code", "Match %"]

    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Header row
    ws.row_dimensions[1].height = 28
    for col, name in enumerate(col_headers, 1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = _HDR_FONT
        cell.fill = _DARK_FILL
        cell.alignment = _CENTER
        cell.border = _THIN

    boq_items: list[dict] = payload.get("boq_items") or []
    current_cat: str | None = None
    row = 2

    for i, item in enumerate(boq_items, 1):
        cat = item.get("section") or item.get("category") or "Uncategorized"
        if cat != current_cat:
            current_cat = cat
            ws.row_dimensions[row].height = 16
            ws.merge_cells(f"A{row}:I{row}")
            hdr = ws.cell(row=row, column=1, value=cat.upper())
            hdr.font = _CAT_FONT
            hdr.fill = _CAT_FILL
            hdr.alignment = _LEFT
            hdr.border = _THIN
            row += 1

        conf = item.get("match_confidence")
        conf_str = f"{conf * 100:.0f}%" if isinstance(conf, float) else "—"

        row_vals: list[Any] = [
            i,
            item.get("description") or item.get("bsr_description") or "—",
            cat,
            item.get("unit") or "—",
            item.get("quantity"),
            item.get("rate"),
            item.get("cost"),
            item.get("bsr_item_no") or "—",
            conf_str,
        ]

        ws.row_dimensions[row].height = 20
        for col, val in enumerate(row_vals, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.border = _THIN
            cell.font = _BODY_FONT
            cell.alignment = _LEFT
            if col in (5, 6, 7) and isinstance(val, (int, float)):
                cell.number_format = "#,##0.00"
                cell.alignment = _RIGHT
            if col == 1:
                cell.alignment = _CENTER
        row += 1

    # Grand total row
    costs: dict = payload.get("costs") or {}
    ws.row_dimensions[row].height = 22
    ws.merge_cells(f"A{row}:F{row}")
    lbl = ws.cell(row=row, column=1, value="GRAND TOTAL  (incl. 5% Contingencies)")
    lbl.font = _TOTAL_FONT
    lbl.fill = _TOTAL_FILL
    lbl.alignment = _RIGHT
    lbl.border = _THIN
    tot = ws.cell(row=row, column=7, value=costs.get("total"))
    tot.font = _TOTAL_FONT
    tot.fill = _TOTAL_FILL
    tot.number_format = "#,##0.00"
    tot.alignment = _RIGHT
    tot.border = _THIN
    for col in (8, 9):
        ws.cell(row=row, column=col).border = _THIN


# ── Sheet 3: Cost Breakdown ───────────────────────────────────────────────────

def _sheet_breakdown(wb: Workbook, payload: dict) -> None:
    ws = wb.create_sheet("Cost Breakdown")
    ws.column_dimensions["A"].width = 36
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 16

    headers = ["Category", "Subtotal (LKR)", "% of Base Total"]
    ws.row_dimensions[1].height = 26
    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = _HDR_FONT
        cell.fill = _DARK_FILL
        cell.alignment = _CENTER
        cell.border = _THIN

    costs: dict = payload.get("costs") or {}
    subtotals: dict[str, float] = costs.get("subtotals") or {}
    base_total = float(costs.get("base_total") or 1)

    sorted_cats = sorted(subtotals.items(), key=lambda x: x[1], reverse=True)
    for r, (cat, amount) in enumerate(sorted_cats, 2):
        ws.row_dimensions[r].height = 20
        pct = amount / base_total
        label = cat.replace("_", " ").title()
        data = [(label, _LEFT, "@"), (amount, _RIGHT, "#,##0.00"), (pct, _RIGHT, "0.00%")]
        for col, (val, align, fmt) in enumerate(data, 1):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = _BODY_FONT
            cell.alignment = align
            cell.border = _THIN
            if fmt != "@":
                cell.number_format = fmt

    # Summary totals
    total_row = len(sorted_cats) + 2
    for offset, (label, key) in enumerate(
        [("Base Total", "base_total"), ("Contingencies (5%)", "contingencies"), ("Grand Total", "total")]
    ):
        r = total_row + offset
        ws.row_dimensions[r].height = 20
        lbl = ws.cell(row=r, column=1, value=label)
        lbl.font = _TOTAL_FONT
        lbl.fill = _TOTAL_FILL
        lbl.border = _THIN
        val = ws.cell(row=r, column=2, value=costs.get(key))
        val.font = _TOTAL_FONT
        val.fill = _TOTAL_FILL
        val.number_format = "#,##0.00"
        val.alignment = _RIGHT
        val.border = _THIN
        ws.cell(row=r, column=3).border = _THIN


# ── Helpers ────────────────────────────────────────────────────────────────────

def _style(
    cell,
    value: Any,
    font=None,
    fill=None,
    align=None,
    height: int | None = None,
    ws=None,
    row: int | None = None,
) -> None:
    cell.value = value
    if font:  cell.font  = font
    if fill:  cell.fill  = fill
    if align: cell.alignment = align
    cell.border = _THIN
    if height and ws and row:
        ws.row_dimensions[row].height = height

